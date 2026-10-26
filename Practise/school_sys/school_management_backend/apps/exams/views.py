import csv, io
from django.http import HttpResponse
from django.db import transaction
from rest_framework import viewsets
from rest_framework.decorators import action
from rest_framework.response import Response
from rest_framework.exceptions import ValidationError
from apps.core.permissions import AdminOrHeadmaster
from apps.students.models import Student
from .models import Subject, Exam, ExamSubject, Result
from .serializers import SubjectSerializer, ExamSerializer, ExamSubjectSerializer, ResultSerializer
class SubjectViewSet(viewsets.ModelViewSet):
    queryset=Subject.objects.all(); serializer_class=SubjectSerializer; permission_classes=[AdminOrHeadmaster]
    filterset_fields=["group","active"]; search_fields=["name","code"]
class ExamViewSet(viewsets.ModelViewSet):
    queryset=Exam.objects.prefetch_related("school_classes","subjects").all(); serializer_class=ExamSerializer; permission_classes=[AdminOrHeadmaster]
    filterset_fields=["academic_year","kind","is_published"]
    @action(detail=True, methods=["post"], url_path="configure-subjects")
    def configure_subjects(self, request, pk=None):
        exam = self.get_object()
        configs = request.data.get("subjects", [])
        mode = request.data.get("mode", "merge").lower()
        delete_existing_results = request.data.get("delete_existing_results", False)
        if not isinstance(configs, list) or not configs:
            raise ValidationError({"subjects": "Expected a non-empty list of subject configurations."})
        if mode not in {"merge", "replace"}:
            raise ValidationError({"mode": "Use 'merge' or 'replace'."})

        prepared = []
        errors = []
        allowed = {"total_marks", "cq_max", "short_max", "mcq_max", "practical_max", "pass_marks",
                   "written_pass_marks", "cq_pass_marks", "mcq_pass_marks", "practical_pass_marks"}
        for item in configs:
            try:
                subject = Subject.objects.get(pk=item.get("subject"))
            except (Subject.DoesNotExist, ValueError, TypeError):
                errors.append({"subject": item.get("subject"), "error": "Unknown subject"})
                continue
            defaults = {key: item[key] for key in allowed if key in item}
            prepared.append((subject, defaults))
        if errors:
            raise ValidationError({"subjects": errors})

        selected_ids = {subject.pk for subject, _ in prepared}
        with transaction.atomic():
            if mode == "replace":
                obsolete = ExamSubject.objects.filter(exam=exam).exclude(subject_id__in=selected_ids)
                has_results = Result.objects.filter(exam_subject__in=obsolete).exists()
                if has_results and not delete_existing_results:
                    raise ValidationError({"mode": "Some removed subjects already have results. Set delete_existing_results=true to explicitly delete those results, or use mode='merge'."})
                obsolete.delete()
            saved = []
            for subject, defaults in prepared:
                cfg, _ = ExamSubject.objects.update_or_create(exam=exam, subject=subject, defaults=defaults)
                saved.append(cfg)
        return Response({"mode": mode, "configured": len(saved), "subjects": ExamSubjectSerializer(saved, many=True).data})

class ResultViewSet(viewsets.ModelViewSet):
    queryset=Result.objects.select_related("student","exam_subject__exam","exam_subject__subject").all()
    serializer_class=ResultSerializer; permission_classes=[AdminOrHeadmaster]
    filterset_fields=["student","exam_subject__exam","exam_subject__subject","passed","is_synthetic"]
    search_fields=["student__name","student__admission_no"]
    @action(detail=False, methods=["post"], url_path="bulk")
    def bulk_upsert(self, request):
        rows=request.data.get("results",[])
        if not isinstance(rows,list) or not rows: raise ValidationError({"results":"Provide a non-empty list."})
        saved=[]
        with transaction.atomic():
            for row in rows:
                try:
                    student=Student.objects.get(pk=row["student"])
                    cfg=ExamSubject.objects.get(pk=row["exam_subject"])
                except (KeyError,Student.DoesNotExist,ExamSubject.DoesNotExist):
                    raise ValidationError({"results":"Each row requires valid student and exam_subject IDs."})
                result,_=Result.objects.update_or_create(student=student,exam_subject=cfg,defaults={k:row[k] for k in ["cq_marks","short_marks","mcq_marks","practical_marks","is_synthetic"] if k in row})
                saved.append(result)
        return Response({"saved":len(saved),"results":self.get_serializer(saved,many=True).data})


    @action(detail=False, methods=["post"], url_path="import-file")
    def import_file(self, request):
        upload=request.FILES.get("file")
        if not upload: raise ValidationError({"file":"Upload CSV or XLSX in multipart field 'file'."})
        filename=upload.name.lower()
        try:
            if filename.endswith(".csv"):
                import csv, io
                records=list(csv.DictReader(io.StringIO(upload.read().decode("utf-8-sig"))))
            elif filename.endswith(".xlsx"):
                from openpyxl import load_workbook
                wb=load_workbook(upload,read_only=True,data_only=True); ws=wb.active
                rows=list(ws.iter_rows(values_only=True))
                records=[dict(zip([str(v).strip() for v in rows[0]],row)) for row in rows[1:] if row and any(v is not None for v in row)]
            else: raise ValidationError({"file":"Only .csv and .xlsx files are supported for result import."})
        except UnicodeDecodeError: raise ValidationError({"file":"CSV must be UTF-8 encoded."})
        saved=0; errors=[]
        for line,row in enumerate(records,start=2):
            try:
                student=Student.objects.get(pk=int(row.get("student_id") or row.get("student")))
                cfg=ExamSubject.objects.get(pk=int(row.get("exam_subject_id") or row.get("exam_subject")))
                values={}
                for field in ["cq_marks","short_marks","mcq_marks","practical_marks"]:
                    raw=row.get(field,0)
                    values[field]=float(raw or 0)
                values["is_synthetic"]=str(row.get("is_synthetic","false")).lower() in ("true","1","yes")
                Result.objects.update_or_create(student=student,exam_subject=cfg,defaults=values); saved+=1
            except Exception as exc: errors.append({"line":line,"error":str(exc)})
        return Response({"saved":saved,"error_count":len(errors),"errors":errors[:100]})
    @action(detail=False, methods=["get"], url_path="export-file")
    def export_file(self, request):
        fmt=request.query_params.get("format","xlsx").lower()
        rows=list(self.filter_queryset(self.get_queryset()))
        headers=["admission_no","student","class","section","exam","year","subject_code","subject","cq","short","mcq","practical","total","grade","grade_point","passed","synthetic"]
        data=[[r.student.admission_no,r.student.name,r.student.school_class.name,r.student.section.name,r.exam_subject.exam.name,r.exam_subject.exam.academic_year,r.exam_subject.subject.code,r.exam_subject.subject.name,float(r.cq_marks),float(r.short_marks),float(r.mcq_marks),float(r.practical_marks),float(r.total_marks),r.grade,float(r.grade_point),r.passed,r.is_synthetic] for r in rows]
        if fmt=="csv":
            response=HttpResponse(content_type="text/csv; charset=utf-8"); response["Content-Disposition"]='attachment; filename="results.csv"'; response.write("\\ufeff")
            w=csv.writer(response); w.writerow(headers); w.writerows(data); return response
        if fmt=="xlsx":
            from openpyxl import Workbook
            from io import BytesIO
            wb=Workbook(); ws=wb.active; ws.title="Results"; ws.append(headers)
            for row in data: ws.append(row)
            buf=BytesIO(); wb.save(buf); buf.seek(0)
            response=HttpResponse(buf.getvalue(),content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            response["Content-Disposition"]='attachment; filename="results.xlsx"'; return response
        if fmt=="docx":
            from io import BytesIO
            from docx import Document
            doc=Document(); doc.add_heading("Exam Results Register",0); table=doc.add_table(rows=1,cols=len(headers)); table.style="Table Grid"
            for cell,val in zip(table.rows[0].cells,headers): cell.text=str(val)
            for row in data:
                cells=table.add_row().cells
                for cell,val in zip(cells,row): cell.text=str(val if val is not None else "")
            buf=BytesIO(); doc.save(buf); buf.seek(0)
            response=HttpResponse(buf.getvalue(),content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
            response["Content-Disposition"]='attachment; filename="results.docx"'; return response
        if fmt=="pdf":
            from io import BytesIO
            from reportlab.lib.pagesizes import landscape,A4
            from reportlab.pdfgen import canvas
            buf=BytesIO(); c=canvas.Canvas(buf,pagesize=landscape(A4)); w,h=landscape(A4)
            c.setFont("Helvetica-Bold",12); c.drawString(25,h-25,"Exam Results Register")
            y=h-48; c.setFont("Helvetica-Bold",7)
            cols=[("Adm",25),("Student",85),("Class",245),("Exam",290),("Subject",385),("Total",525),("Grade",570),("GP",610)]
            for label,x in cols:c.drawString(x,y,label)
            y-=13;c.setFont("Helvetica",7)
            for row in data:
                if y<28:c.showPage();y=h-30;c.setFont("Helvetica",7)
                vals=[str(row[0])[:9],str(row[1])[:30],str(row[2]),str(row[4])[:18],str(row[7])[:24],str(row[12]),str(row[13]),str(row[14])]
                for val,(_,x) in zip(vals,cols): c.drawString(x,y,val)
                y-=11
            c.save();buf.seek(0);response=HttpResponse(buf.getvalue(),content_type="application/pdf")
            response["Content-Disposition"]='attachment; filename="results.pdf"';return response
        raise ValidationError({"format":"Supported formats: csv, xlsx, docx, pdf."})
    @action(detail=False, methods=["get"], url_path="export-csv")
    def export_csv(self, request):
        qs=self.filter_queryset(self.get_queryset())
        response=HttpResponse(content_type="text/csv; charset=utf-8"); response["Content-Disposition"]='attachment; filename="results.csv"'; response.write("\ufeff")
        w=csv.writer(response); w.writerow(["student_admission_no","student_name","class","section","exam","year","subject_code","subject","cq","short","mcq","practical","total","grade","grade_point","passed","synthetic"])
        for r in qs:
            w.writerow([r.student.admission_no,r.student.name,r.student.school_class.name,r.student.section.name,r.exam_subject.exam.name,r.exam_subject.exam.academic_year,r.exam_subject.subject.code,r.exam_subject.subject.name,r.cq_marks,r.short_marks,r.mcq_marks,r.practical_marks,r.total_marks,r.grade,r.grade_point,r.passed,r.is_synthetic])
        return response
