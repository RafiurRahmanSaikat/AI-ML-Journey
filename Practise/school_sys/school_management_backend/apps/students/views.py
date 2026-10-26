import csv, io
from datetime import date
from django.db import transaction
from django.http import HttpResponse
from rest_framework import viewsets, status
from rest_framework.decorators import action
from rest_framework.response import Response
from rest_framework.exceptions import ValidationError
from apps.core.permissions import AdminOrHeadmaster
from apps.academics.models import SchoolClass, Section
from .models import Student
from .serializers import StudentSerializer
class StudentViewSet(viewsets.ModelViewSet):
    queryset=Student.objects.select_related("school_class","section").all()
    serializer_class=StudentSerializer
    permission_classes=[AdminOrHeadmaster]
    filterset_fields=["school_class","section","group","active"]
    search_fields=["name","admission_no","father_name","phone"]
    ordering_fields=["roll","name","created_at"]
    @action(detail=False, methods=["post"], url_path="bulk")
    def bulk_create(self, request):
        rows=request.data.get("students",[])
        if not isinstance(rows,list) or not rows: raise ValidationError({"students":"Provide a non-empty list of student objects."})
        serializer=self.get_serializer(data=rows,many=True)
        serializer.is_valid(raise_exception=True)
        created=[]
        with transaction.atomic():
            for item in serializer.validated_data:
                obj=Student.objects.create(**item); created.append(obj)
        return Response({"created":len(created),"students":self.get_serializer(created,many=True).data},status=status.HTTP_201_CREATED)
    @action(detail=False, methods=["get"], url_path="export")
    def export_csv(self, request):
        response=HttpResponse(content_type="text/csv; charset=utf-8")
        response["Content-Disposition"]='attachment; filename="students.csv"'
        response.write("\ufeff")
        writer=csv.writer(response)
        fields=["admission_no","roll","name","father_name","mother_name","school_class","section","group","date_of_birth","phone","address","active"]
        writer.writerow(fields)
        for s in self.filter_queryset(self.get_queryset()):
            writer.writerow([s.admission_no,s.roll,s.name,s.father_name,s.mother_name,s.school_class.name,s.section.name,s.group,s.date_of_birth or "",s.phone,s.address,s.active])
        return response

    @action(detail=False, methods=["get"], url_path="export-file")
    def export_file(self, request):
        fmt=request.query_params.get("format","xlsx").lower()
        rows=list(self.filter_queryset(self.get_queryset()))
        headers=["admission_no","roll","name","father_name","mother_name","class","section","group","date_of_birth","phone","address","active"]
        data=[[s.admission_no,s.roll,s.name,s.father_name,s.mother_name,s.school_class.name,s.section.name,s.group,str(s.date_of_birth or ""),s.phone,s.address,s.active] for s in rows]
        if fmt=="csv":
            response=HttpResponse(content_type="text/csv; charset=utf-8"); response["Content-Disposition"]='attachment; filename="students.csv"'; response.write("\\ufeff")
            w=csv.writer(response); w.writerow(headers); w.writerows(data); return response
        if fmt=="xlsx":
            from openpyxl import Workbook
            from io import BytesIO
            wb=Workbook(); ws=wb.active; ws.title="Students"; ws.append(headers)
            for row in data: ws.append(row)
            buf=BytesIO(); wb.save(buf); buf.seek(0)
            response=HttpResponse(buf.getvalue(),content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            response["Content-Disposition"]='attachment; filename="students.xlsx"'; return response
        if fmt=="docx":
            from io import BytesIO
            from docx import Document
            doc=Document(); doc.add_heading("Student Register",0); table=doc.add_table(rows=1, cols=len(headers)); table.style="Table Grid"
            for cell,val in zip(table.rows[0].cells,headers): cell.text=str(val)
            for row in data:
                cells=table.add_row().cells
                for cell,val in zip(cells,row): cell.text=str(val if val is not None else "")
            buf=BytesIO(); doc.save(buf); buf.seek(0)
            response=HttpResponse(buf.getvalue(),content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
            response["Content-Disposition"]='attachment; filename="students.docx"'; return response
        if fmt=="pdf":
            from io import BytesIO
            from reportlab.lib.pagesizes import landscape, A4
            from reportlab.pdfgen import canvas
            buf=BytesIO(); c=canvas.Canvas(buf,pagesize=landscape(A4)); w,h=landscape(A4)
            c.setFont("Helvetica-Bold",14); c.drawString(30,h-30,"Student Register")
            y=h-55; c.setFont("Helvetica-Bold",7); c.drawString(25,y,"Adm No"); c.drawString(90,y,"Roll"); c.drawString(120,y,"Name"); c.drawString(300,y,"Class"); c.drawString(350,y,"Sec"); c.drawString(390,y,"Group"); c.drawString(470,y,"Phone")
            y-=15; c.setFont("Helvetica",7)
            for row in data:
                if y<35: c.showPage(); y=h-35; c.setFont("Helvetica",7)
                c.drawString(25,y,str(row[0])[:12]); c.drawString(90,y,str(row[1] or "-")); c.drawString(120,y,str(row[2])[:32]); c.drawString(300,y,str(row[5])); c.drawString(350,y,str(row[6])); c.drawString(390,y,str(row[7])); c.drawString(470,y,str(row[9])[:18]); y-=12
            c.save(); buf.seek(0); response=HttpResponse(buf.getvalue(),content_type="application/pdf")
            response["Content-Disposition"]='attachment; filename="students.pdf"'; return response
        raise ValidationError({"format":"Supported formats: csv, xlsx, docx, pdf."})
    @action(detail=False, methods=["post"], url_path="import-file")
    def import_csv(self, request):
        upload=request.FILES.get("file")
        if not upload: raise ValidationError({"file":"Upload a CSV file using multipart/form-data field 'file'."})
        try:
            if upload.name.lower().endswith(".xlsx"):
                from openpyxl import load_workbook
                wb=load_workbook(upload,read_only=True,data_only=True); ws=wb.active
                sheet=list(ws.iter_rows(values_only=True))
                if not sheet: raise ValidationError({"file":"Workbook is empty."})
                headers=[str(v).strip() if v is not None else "" for v in sheet[0]]
                rows=[dict(zip(headers,row)) for row in sheet[1:] if row and any(v is not None for v in row)]
                reader=rows
            else:
                text=upload.read().decode("utf-8-sig")
                reader=list(csv.DictReader(io.StringIO(text)))
        except UnicodeDecodeError: raise ValidationError({"file":"CSV must be UTF-8 encoded."})
        created=0; updated=0; errors=[]
        for line,row in enumerate(reader,start=2):
            try:
                cls=SchoolClass.objects.get(name__iexact=(row.get("school_class") or row.get("Class") or row.get("class") or "").strip())
                sec=Section.objects.get(school_class=cls,name__iexact=(row.get("section") or row.get("Section") or "A").strip())
                dob=(row.get("date_of_birth") or row.get("Date of Birth") or "").strip()
                if dob in ("-","0001-01-01","0020-09-08","0013-12-30"): dob=""
                payload={"roll":int(row["roll"]) if (row.get("roll") or row.get("Roll") or "").strip().isdigit() else None,
                    "name":(row.get("name") or row.get("Name") or "").strip(),
                    "father_name":(row.get("father_name") or row.get("Father Name") or "").strip(),
                    "mother_name":(row.get("mother_name") or row.get("Mother Name") or "").strip(),
                    "school_class":cls,"section":sec,"group":(row.get("group") or "HUMANITIES").upper(),
                    "date_of_birth":date.fromisoformat(dob) if dob else None,
                    "phone":(row.get("phone") or row.get("Phone") or "").strip().replace('="','').replace('"',''),
                    "address":(row.get("address") or row.get("Address") or "").strip()}
                if not payload["name"]: raise ValueError("Name is required")
                adm=(row.get("admission_no") or row.get("ID") or "").strip()
                if adm and Student.objects.filter(admission_no=adm).exists():
                    obj=Student.objects.get(admission_no=adm)
                    for k,v in payload.items(): setattr(obj,k,v)
                    obj.save(); updated+=1
                else:
                    Student.objects.create(admission_no=adm or "",**payload); created+=1
            except Exception as exc:
                errors.append({"line":line,"error":str(exc)})
        return Response({"created":created,"updated":updated,"error_count":len(errors),"errors":errors[:100]},status=200)
